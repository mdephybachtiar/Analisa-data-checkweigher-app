import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.chart import LineChart, Reference
from datetime import datetime
import io

# --- KONFIGURASI HALAMAN MOBILE-FRIENDLY ---
st.set_page_config(page_title="Checkweigher Analysis App", page_icon="📊", layout="centered")

st.title("📊 Aplikasi Analisa Checkweigher")
st.write("Upload file CSV dari mesin, sesuaikan parameter, dan download laporan Excel Anda langsung dari HP.")

# --- SIDEBAR / FORM INPUT PARAMETER ---
st.sidebar.header("⚙️ Pengaturan Parameter")

customer_name = st.sidebar.text_input("Nama Customer", value="PT. Centra Proteina Prima")
line_produksi = st.sidebar.text_input("Line Produksi", value="Line 7")
nomor_mesin = st.sidebar.text_input("Nomor Mesin", value="04")
performance_by = st.sidebar.text_input("Operator / PIC", value="Dephy Bachtiar")

st.sidebar.markdown("---")
berat_kemasan = st.sidebar.number_input("Berat Kemasan (g)", value=15)
batas_bawah = st.sidebar.number_input("Batas Bawah Error (g)", value=750)
batas_atas = st.sidebar.number_input("Batas Atas Error (g)", value=850)
reject_under = st.sidebar.number_input("Reject Under Limit", value=813)
reject_over = st.sidebar.number_input("Reject Over Limit", value=830)

# --- UTAMA: UPLOAD FILE CSV ---
st.markdown("### 📁 Upload Data Mentah")
uploaded_file = st.file_uploader("Pilih file CSV data checkweigher", type=["csv"])

def evaluate_status(weight):
    if weight < reject_under: return 'Under'
    elif weight > reject_over: return 'Over'
    else: return 'OK'

if uploaded_file is not None:
    st.success("File berhasil di-upload!")
    
    if st.button("🚀 Proses & Generate Laporan Excel", type="primary"):
        try:
            # 1. Baca CSV dari file yang di-upload
            df = pd.read_csv(uploaded_file)
            df.columns = [col.replace('="', '').replace('"', '') if isinstance(col, str) else col for col in df.columns]
            for col in ['Time', 'Weight(g)', 'Status']:
                if col in df.columns:
                    df[col] = df[col].astype(str).str.replace('="', '').str.replace('"', '')
            df['Weight(g)'] = pd.to_numeric(df['Weight(g)'], errors='coerce')

            # 2. Filter Validasi
            df_valid = df[(df['Weight(g)'] >= batas_bawah) & (df['Weight(g)'] <= batas_atas)].copy()
            
            if not df_valid.empty:
                df_valid['Net_Weight(g)'] = df_valid['Weight(g)'] - berat_kemasan
                df_valid['Evaluated_Status'] = df_valid['Weight(g)'].apply(evaluate_status)
                df_valid['Time'] = pd.to_datetime(df_valid['Time'])
                
                tanggal_data_obj = df_valid['Time'].iloc[0]
                tanggal_str_file = tanggal_data_obj.strftime("%d_%b_%Y") 
                tanggal_format_d4 = tanggal_data_obj.strftime("%d-%m-%Y") 
                
                output_filename = f'Laporan_Checkweigher_{tanggal_str_file}.xlsx'
                
                df_valid['Hour_Only'] = df_valid['Time'].dt.strftime('%H:00')

                summary_per_jam = df_valid.groupby('Hour_Only').agg(
                    Total_Produk=('Weight(g)', 'count'),
                    Rata_Bruto=('Weight(g)', 'mean'),
                    Rata_Netto=('Net_Weight(g)', 'mean'), 
                    Deviasi=('Net_Weight(g)', 'std')
                ).reset_index()

                reject_u = df_valid[df_valid['Evaluated_Status'] == 'Under'].groupby('Hour_Only').size().reset_index(name='Under')
                reject_o = df_valid[df_valid['Evaluated_Status'] == 'Over'].groupby('Hour_Only').size().reset_index(name='Over')

                hasil = pd.merge(summary_per_jam, reject_u, on='Hour_Only', how='left').fillna(0)
                hasil = pd.merge(hasil, reject_o, on='Hour_Only', how='left').fillna(0)
                hasil['Rejection_Rate'] = ((hasil['Under'] + hasil['Over']) / hasil['Total_Produk']) * 100

                # 3. Masukkan ke Template Excel
                template_file = 'Template New.xlsx'
                if not openpyxl.os.path.exists(template_file):
                    st.error("❌ File 'Template New.xlsx' tidak ditemukan di folder server aplikasi!")
                else:
                    wb = openpyxl.load_workbook(template_file)
                    sheet = wb['Analisa data']
                    
                    # ISI IDENTITAS
                    sheet['D3'] = customer_name
                    sheet['D4'] = tanggal_format_d4
                    sheet['D5'] = line_produksi
                    sheet['D6'] = nomor_mesin
                    sheet['D7'] = performance_by
                    for cell_id in ['D3', 'D4', 'D5', 'D6', 'D7']:
                        sheet[cell_id].alignment = Alignment(horizontal='center', vertical='center')

                    # ISI SUMMARY ATAS
                    total_semua = hasil['Total_Produk'].sum()
                    sheet['I3'] = total_semua
                    sheet['I4'] = round(df_valid['Weight(g)'].mean(), 2)       
                    sheet['I5'] = round(df_valid['Net_Weight(g)'].mean(), 2)   
                    total_rej = hasil['Under'].sum() + hasil['Over'].sum()
                    sheet['I6'] = f"{round((total_rej / total_semua) * 100, 2)}%" if total_semua > 0 else "0%"
                    sheet['I7'] = round(df_valid['Net_Weight(g)'].std(), 2)
                    sheet['I2'] = ""

                    # ISI TABEL BAWAH
                    start_row = 11
                    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

                    for index, row in hasil.iterrows():
                        r = start_row + index
                        sheet[f'B{r}'], sheet[f'C{r}'] = row['Hour_Only'], row['Total_Produk']
                        sheet[f'D{r}'], sheet[f'E{r}'] = round(row['Rata_Bruto'], 2), round(row['Rata_Netto'], 2)
                        sheet[f'F{r}'] = round(row['Deviasi'], 2) if pd.notna(row['Deviasi']) else 0
                        sheet[f'G{r}'], sheet[f'H{r}'] = row['Under'], row['Over']
                        sheet[f'I{r}'] = f"{round(row['Rejection_Rate'], 2)}%"
                        for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
                            sheet[f'{col}{r}'].border = thin_border
                            sheet[f'{col}{r}'].alignment = Alignment(horizontal='center', vertical='center')

                    # DASHBOARD GRAFIK
                    if 'Dashboard Grafik' in wb.sheetnames: del wb['Dashboard Grafik']
                    chart_sheet = wb.create_sheet(title="Dashboard Grafik")
                    
                    chart_sheet['A1'], chart_sheet['B1'], chart_sheet['C1'] = "Jam", "Reject Rate (%)", "Deviasi (g)"
                    for idx, r in hasil.iterrows():
                        chart_sheet.cell(row=idx+2, column=1, value=r['Hour_Only'])
                        chart_sheet.cell(row=idx+2, column=2, value=round(r['Rejection_Rate'], 2))
                        chart_sheet.cell(row=idx+2, column=3, value=round(r['Deviasi'], 2) if pd.notna(r['Deviasi']) else 0)

                    cats_ref = Reference(chart_sheet, min_col=1, min_row=2, max_row=len(hasil)+1)
                    
                    c1 = LineChart()
                    c1.title = f"Trend Rejection Rate (%) - {tanggal_str_file}"
                    c1.y_axis.title = "Persentase (%)"
                    c1.add_data(Reference(chart_sheet, min_col=2, min_row=1, max_row=len(hasil)+1), titles_from_data=True)
                    c1.set_categories(cats_ref)
                    chart_sheet.add_chart(c1, "G2")

                    c2 = LineChart()
                    c2.title = f"Trend Deviasi (g) - {tanggal_str_file}"
                    c2.y_axis.title = "Gram (g)"
                    c2.add_data(Reference(chart_sheet, min_col=3, min_row=1, max_row=len(hasil)+1), titles_from_data=True)
                    c2.set_categories(cats_ref)
                    chart_sheet.add_chart(c2, "G17")

                    # Simpan ke buffer memori agar bisa langsung didownload
                    virtual_workbook = io.BytesIO()
                    wb.save(virtual_workbook)
                    virtual_workbook.seek(0)
                    
                    st.success("✨ Analisis Selesai! Silakan download file hasil di bawah.")
                    
                    # Tombol Download untuk HP
                    st.download_button(
                        label="📥 Download Laporan Excel",
                        data=virtual_workbook,
                        file_name=output_filename,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            else:
                st.error("❌ ERROR: Semua data berada di luar batas error atau tidak valid.")
        except Exception as e:
            st.error(f"❌ Terjadi kesalahan saat memproses data: {e}")
